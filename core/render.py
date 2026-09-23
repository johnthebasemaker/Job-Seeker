"""Render the canonical markdown to an ATS-safe PDF and DOCX.

Single column, real text, no tables, no images, no text boxes - the four
things that make resume parsers drop content. The contact block is added
here, which is the only place it meets the resume body.

The PDF goes through Typst, which ships as a pip wheel and needs no system
packages, so it works on Streamlit Community Cloud where LaTeX would not.
"""

from __future__ import annotations

import io
import re
from typing import Any

from .resume import Section, parse

FONT_STACK = '("Libertinus Serif", "Linux Libertine", "New Computer Modern", "DejaVu Serif")'


class RenderError(RuntimeError):
    pass


# ------------------------------------------------------------------ shared
def contact_lines(contact: dict[str, str]) -> list[str]:
    order = ["email", "phone", "location", "linkedin", "github", "website"]
    lines = []
    for key in order:
        value = str(contact.get(key) or "").strip()
        if value:
            lines.append(re.sub(r"^https?://", "", value))
    return lines


def display_name(contact: dict[str, str]) -> str:
    return str(contact.get("full_name") or "Your Name").strip()


def safe_filename(contact: dict[str, str], job: dict[str, Any] | None, ext: str) -> str:
    parts = [display_name(contact)]
    if job:
        parts += [job.get("company") or "", job.get("title") or ""]
    slug = re.sub(r"[^A-Za-z0-9]+", "_", " ".join(p for p in parts if p)).strip("_")
    return f"{slug[:80] or 'resume'}.{ext}"


# --------------------------------------------------------------------- PDF
def _typst_string(value: str) -> str:
    """A Typst string literal. Nothing inside is ever parsed as markup."""
    escaped = (
        (value or "")
        .replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\n", " ")
        .replace("\r", " ")
    )
    return f'"{escaped}"'


def _typst_array(values: list[str]) -> str:
    if not values:
        return "()"
    inner = ", ".join(_typst_string(v) for v in values)
    return f"({inner},)"


def _typst_data(sections: list[Section], contact: dict[str, str]) -> str:
    blocks = []
    for section in sections:
        entries = ", ".join(
            "(heading: {h}, meta: {m}, body: {b}, bullets: {bl})".format(
                h=_typst_string(entry.heading),
                m=_typst_string(entry.meta),
                b=_typst_string(entry.body),
                bl=_typst_array(entry.bullets),
            )
            for entry in section.entries
        )
        blocks.append(
            "(title: {t}, body: {b}, bullets: {bl}, entries: ({e}{comma}))".format(
                t=_typst_string(section.title),
                b=_typst_string(section.body),
                bl=_typst_array(section.bullets),
                e=entries,
                comma="," if entries else "",
            )
        )
    sections_src = ", ".join(blocks)
    return (
        "#let data = (name: {name}, contact: {contact}, sections: ({sections}{comma}))"
        .format(
            name=_typst_string(display_name(contact)),
            contact=_typst_array(contact_lines(contact)),
            sections=sections_src,
            comma="," if blocks else "",
        )
    )


TYPST_TEMPLATE = """
#set document(title: {title}, author: {author})
#set page(paper: "a4", margin: (x: 1.5cm, y: 1.3cm))
#set text(font: %s, size: 10.5pt, hyphenate: false)
#set par(justify: false, leading: 0.60em, spacing: 0.62em)

#align(center)[
  #text(size: 18pt, weight: "bold")[#data.name]
  #linebreak()
  #text(size: 9.5pt)[#data.contact.join("  |  ")]
]
#v(2pt)

#for sec in data.sections [
  #block(above: 13pt, below: 6pt, stack(
    spacing: 3pt,
    text(size: 10.5pt, weight: "bold", tracking: 0.6pt, upper(sec.title)),
    line(length: 100%%, stroke: 0.6pt + luma(140)),
  ))
  #if sec.body != "" [ #sec.body ]
  #for b in sec.bullets [
    #par(hanging-indent: 11pt)[#sym.bullet #h(4pt) #b]
  ]
  #for e in sec.entries [
    #block(above: 8pt, below: 3pt)[
      #text(weight: "bold")[#e.heading]
      #if e.meta != "" [ #linebreak() #text(size: 9.5pt, fill: luma(70))[#e.meta] ]
      #if e.body != "" [ #linebreak() #e.body ]
    ]
    #for b in e.bullets [
      #par(hanging-indent: 11pt)[#sym.bullet #h(4pt) #b]
    ]
  ]
]
""" % FONT_STACK


def typst_source(markdown: str, contact: dict[str, str]) -> str:
    sections = parse(markdown)
    if not sections:
        raise RenderError("The resume has no '## Section' headings yet.")
    header = _typst_data(sections, contact)
    body = TYPST_TEMPLATE.format(
        title=_typst_string(f"{display_name(contact)} - Resume"),
        author=_typst_string(display_name(contact)),
    )
    return f"{header}\n{body}"


def resume_pdf(markdown: str, contact: dict[str, str]) -> bytes:
    try:
        import typst
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise RenderError("The 'typst' package is not installed") from exc

    source = typst_source(markdown, contact)
    try:
        return typst.compile(source.encode("utf-8"))
    except Exception as exc:
        raise RenderError(f"Typst could not build the PDF: {exc}") from exc


# -------------------------------------------------------------------- DOCX
def resume_docx(markdown: str, contact: dict[str, str]) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise RenderError("python-docx is not installed") from exc

    document = Document()
    for style_name in ("Normal",):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(display_name(contact))
    run.bold = True
    run.font.size = Pt(18)

    lines = contact_lines(contact)
    if lines:
        details = document.add_paragraph()
        details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        detail_run = details.add_run("  |  ".join(lines))
        detail_run.font.size = Pt(9.5)

    for section in parse(markdown):
        heading = document.add_heading(section.title.upper(), level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            run.font.size = Pt(11)
        if section.body:
            document.add_paragraph(section.body)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        for entry in section.entries:
            paragraph = document.add_paragraph()
            entry_run = paragraph.add_run(entry.heading)
            entry_run.bold = True
            if entry.meta:
                meta = document.add_paragraph()
                meta_run = meta.add_run(entry.meta)
                meta_run.italic = True
                meta_run.font.size = Pt(9.5)
            if entry.body:
                document.add_paragraph(entry.body)
            for bullet in entry.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
